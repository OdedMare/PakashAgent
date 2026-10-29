"""The import routes: a preview that writes nothing, then a confirm (D7/D39)."""

import io
import json

from openpyxl import Workbook

from app.common.sessions import ROLE_MEMBER
from tests.test_schedule_api import MORNING, _build_app, _client


def _xlsx(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


WEEK = _xlsx([["", "2/11/26", "3/11/26"], [MORNING, "דנה", "יוסי ב."]])


def _preview(client, content=WEEK, options=None):
    return client.post(
        "/api/schedule/import/preview",
        files={"file": ("week.xlsx", content, "application/octet-stream")},
        data={"options": json.dumps(options or {})},
    )


def test_a_preview_reads_the_file_and_stores_nothing():
    app, repo = _build_app()
    response = _preview(_client(app))
    assert response.status_code == 200, response.text
    body = response.json()
    assert [(row["employee"], row["date"]) for row in body["assignments"]] == [
        ("דנה", "2026-11-02"), ("יוסי", "2026-11-03")]
    assert body["sheet"][1][1] == "דנה"
    assert repo.schedules == {} and repo.model_calls == []


def test_a_member_cannot_preview_or_confirm():
    app, _repo = _build_app()
    client = _client(app, role=ROLE_MEMBER)
    assert _preview(client).status_code in (401, 403)
    assert client.post("/api/schedule/import/confirm",
                       json={"assignments": []}).status_code in (401, 403)


def test_options_re_read_the_file_with_the_managers_corrections():
    app, _repo = _build_app()
    body = _preview(_client(app), options={"names": {"יוסי ב.": ""}}).json()
    assert [row["employee"] for row in body["assignments"]] == ["דנה"]


def test_broken_options_are_a_hebrew_error():
    app, _repo = _build_app()
    response = _client(app).post(
        "/api/schedule/import/preview",
        files={"file": ("week.xlsx", WEEK, "application/octet-stream")},
        data={"options": "{not json"})
    assert response.status_code == 400
    assert "אפשרויות" in response.text


def test_confirming_opens_a_period_with_imported_rows():
    app, repo = _build_app()
    client = _client(app)
    body = _preview(client).json()
    response = client.post("/api/schedule/import/confirm", json={
        "assignments": body["assignments"], "filename": "week.xlsx",
        "unavailability": [{"employee": "דנה", "date": "2026-11-04", "reason": "חופש"}],
    })
    assert response.status_code == 200, response.text
    schedule = response.json()
    assert (schedule["starts_on"], schedule["ends_on"]) == ("2026-11-02", "2026-11-03")
    assert {row["source"] for row in schedule["assignments"]} == {"imported"}
    assert all("week.xlsx" in row["reason"] for row in schedule["assignments"])
    assert repo.availability_rows[0]["employee"] == "דנה"
    assert repo.changes[-1]["action"] == "imported"


def test_an_overlapping_period_is_named_in_the_preview_and_refused_on_confirm():
    app, _repo = _build_app()
    client = _client(app)
    client.post("/api/schedule/blank", json={"starts_on": "2026-11-01", "ends_on": "2026-11-07"})
    body = _preview(client).json()
    assert body["overlaps"][0]["starts_on"] == "2026-11-01"
    refused = client.post("/api/schedule/import/confirm", json={"assignments": body["assignments"]})
    assert refused.status_code == 400 and "כבר קיים סידור" in refused.text


def test_merging_into_a_draft_rewrites_only_the_imported_dates():
    app, repo = _build_app()
    client = _client(app)
    draft = client.post("/api/schedule/blank",
                        json={"starts_on": "2026-11-01", "ends_on": "2026-11-07"}).json()
    client.post("/api/schedule/assign", json={
        "schedule_id": draft["id"], "shift_name": MORNING, "slot_date": "2026-11-05",
        "employee": "יוסי"})
    body = _preview(client).json()
    response = client.post("/api/schedule/import/confirm", json={
        "assignments": body["assignments"], "schedule_id": draft["id"]})
    assert response.status_code == 200, response.text
    rows = {(row["employee"], row["date"], row["source"]) for row in response.json()["assignments"]}
    assert rows == {("דנה", "2026-11-02", "imported"), ("יוסי", "2026-11-03", "imported"),
                    ("יוסי", "2026-11-05", "manager")}


def test_rows_outside_the_chosen_draft_are_refused():
    app, _repo = _build_app()
    client = _client(app)
    draft = client.post("/api/schedule/blank",
                        json={"starts_on": "2026-11-03", "ends_on": "2026-11-07"}).json()
    body = _preview(client).json()
    response = client.post("/api/schedule/import/confirm", json={
        "assignments": body["assignments"], "schedule_id": draft["id"]})
    assert response.status_code == 400 and "חורגים" in response.text


def test_a_workbook_with_a_tab_per_month_opens_on_the_one_covering_today():
    from app.common.time_context import israel_today

    today = israel_today()
    book = Workbook()
    book.active.title = "ישן"
    book.active.append(["", "2/1/20", "3/1/20"])
    book.active.append([MORNING, "דנה", "יוסי"])
    current = book.create_sheet("החודש")
    current.append(["", today.strftime("%d/%m/%Y")])
    current.append([MORNING, "דנה"])
    buffer = io.BytesIO()
    book.save(buffer)
    app, _repo = _build_app()
    body = _preview(_client(app), content=buffer.getvalue()).json()
    assert body["sheets"] == ["ישן", "החודש"] and body["sheet_name"] == "החודש"
